"""Exercise the real WebView2 window against isolated Mock/E5 demo data.

Run from the repository root with python -m agent_workbench.scripts.verify_desktop.
This verifies DOM interactions inside the native shell, not the OS file picker.
"""
from __future__ import annotations

import base64
import json
import io
import os
import socket
import sqlite3
import tempfile
import threading
import time
import traceback
from pathlib import Path

import pymupdf
import webview

from agent_workbench.app import find_available_port, run_desktop
from agent_workbench.bootstrap import build_runtime
from agent_workbench.core.models import ProviderChunk
from agent_workbench.providers.mock import MockProvider


class CompactionSmokeProvider(MockProvider):
    """Explicit fixture: tests native plumbing, not real-model summary quality."""
    supports_semantic_compaction = True

    async def stream(self, messages, tools, cancel_event):
        if 'conversation compaction' in messages[0]['content']:
            yield ProviderChunk(type='content', delta='## Goal\nVerify native desktop.\n## Constraints\nKeep original history.\n## Pending\nReview upgrade evidence.')
            yield ProviderChunk(type='usage', usage={'prompt_tokens':100,'completion_tokens':20,'total_tokens':120})
            yield ProviderChunk(type='finish', finish_reason='stop')
            return
        async for chunk in super().stream(messages, tools, cancel_event):
            yield chunk


def wait_for(check, *, timeout=90):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        value = check()
        if value:
            return value
        time.sleep(0.1)
    raise TimeoutError('Native desktop verification condition timed out')


def main():
    temp = Path(__file__).resolve().parents[1] / '.tmp'
    temp.mkdir(exist_ok=True)
    os.environ['TEMP'] = os.environ['TMP'] = str(temp)
    root = Path(tempfile.mkdtemp(prefix='desktop-v15-', dir=temp))
    runtime = build_runtime(data_dir=root / 'data', environment={})
    workspace = root / 'workspace'
    workspace.mkdir()
    target = workspace / 'sample.txt'
    target.write_text('original evidence\n', encoding='utf-8')
    runtime.set_workspace(workspace)
    port = find_available_port()
    report = {'status': 'running', 'data_dir': str(root), 'checks': [], 'port': port}

    def exercise():
        window = None
        stage = 'window_loaded'
        try:
            window = wait_for(lambda: webview.windows and webview.windows[0])
            if not window.events.loaded.wait(60):
                raise TimeoutError('WebView2 page did not load')
            js = window.evaluate_js
            wait_for(lambda: js('!!window.pywebview?.api?.choose_workspace'))
            wait_for(lambda: js('document.querySelector("#fileTree").textContent.includes("sample.txt")'))
            js('window.__awbErrors = []; window.addEventListener("error", e => window.__awbErrors.push(e.message)); window.addEventListener("unhandledrejection", e => window.__awbErrors.push(String(e.reason)));')
            report['renderer'] = webview.renderer
            report['checks'].append('native_window_and_bridge_ready')
            wait_for(lambda: js('document.querySelector(".mobius-stage canvas")?.dataset.animating === "true"'))
            report['checks'].append('native_mobius_surface_renders_and_animates')
            stage = 'native_tray_close_choice'
            from System import Action
            from System.Drawing import Bitmap, Rectangle
            from System.Drawing.Imaging import ImageFormat
            wait_for(lambda: hasattr(window, '_desktop_tray'))
            tray = window._desktop_tray
            def native_call(callback):
                window.native.Invoke(Action(callback))
            import ctypes
            def caption_color(form):
                from agent_workbench.desktop_theme import apply_light_caption
                results = apply_light_caption(form)
                assert results[35] == 0, results
                return 0xFFFFFF
            assert caption_color(window.native) == 0xFFFFFF
            report['checks'].append('native_light_caption_matches_workbench')
            def begin_close():
                window.native.BeginInvoke(Action(lambda: window.native.Close()))
                wait_for(lambda: tray.dialog is not None and tray.dialog.Visible)
            begin_close()
            assert caption_color(tray.dialog) == 0xFFFFFF
            def capture_choice():
                dialog = tray.dialog
                image = Bitmap(dialog.Width, dialog.Height)
                try:
                    dialog.DrawToBitmap(image, Rectangle(0, 0, dialog.Width, dialog.Height))
                    image.Save(str(root / 'close-choice.png'), ImageFormat.Png)
                finally:
                    image.Dispose()
            native_call(capture_choice)
            native_call(lambda: tray.dialog.CancelButton.PerformClick())
            wait_for(lambda: tray.dialog is None)
            assert window.native.Visible and tray.controller.preference.load() == 'ask'
            report['checks'].append('native_first_close_cancel_preserves_window')
            from System.Windows.Forms import Application
            window.native.BeginInvoke(Action(lambda: tray.error(OSError('Fixture detail: no real user paths'))))
            def error_dialog():
                return next((form for form in Application.OpenForms if form != window.native and form.Text == 'AgentWorkbench'), None)
            wait_for(lambda: error_dialog() is not None)
            def check_error():
                dialog = error_dialog()
                assert caption_color(dialog) == 0xFFFFFF
                details = next(control for control in dialog.Controls if control.GetType().Name == 'TextBox')
                assert not details.Visible
                next(control for control in dialog.Controls if control.Text == '详细信息').PerformClick()
                assert details.Visible and 'Fixture detail' in details.Text
                image = Bitmap(dialog.Width, dialog.Height)
                try:
                    dialog.DrawToBitmap(image, Rectangle(0, 0, dialog.Width, dialog.Height))
                    image.Save(str(root / 'error-details.png'), ImageFormat.Png)
                finally: image.Dispose()
                dialog.AcceptButton.PerformClick()
            native_call(check_error)
            report['checks'].append('native_error_details_collapsed_and_expandable')
            begin_close()
            native_call(lambda: next(c for c in tray.dialog.Controls if c.Text == '隐藏到托盘').PerformClick())
            wait_for(lambda: not window.native.Visible)
            assert tray.tray.Visible and tray.controller.preference.load() == 'hide'
            wait_for(lambda: js('document.querySelector(".mobius-stage canvas")?.dataset.animating === "false" && document.body.classList.contains("ambient-suspended")'))
            from agent_workbench.app import wait_for_health
            assert wait_for_health(port)['status'] == 'ok'
            native_call(lambda: tray.open_item.PerformClick())
            wait_for(lambda: window.native.Visible)
            wait_for(lambda: js('document.querySelector(".mobius-stage canvas")?.dataset.animating === "true"'))
            from System.Windows.Forms import FormWindowState
            native_call(lambda: setattr(window.native, 'WindowState', FormWindowState.Minimized))
            wait_for(lambda: js('document.querySelector(".mobius-stage canvas")?.dataset.animating === "false"'))
            native_call(lambda: tray.open_item.PerformClick())
            wait_for(lambda: js('document.querySelector(".mobius-stage canvas")?.dataset.animating === "true"'))
            report['checks'].append('native_mobius_pauses_in_tray_and_minimized_and_resumes')
            native_call(window.native.Close)
            assert not window.native.Visible and tray.dialog is None
            native_call(lambda: tray.reset_item.PerformClick())
            assert window.native.Visible and tray.controller.preference.load() == 'ask'
            report['checks'].append('native_tray_hide_restore_remember_reset_keeps_backend_alive')
            assert js('document.querySelector(".topbar #contextStatus") !== null && getComputedStyle(document.querySelector(".sidebar")).borderRightWidth === "0px"')
            assert js('document.querySelector(".ambient-plane,#dynamicBackground") === null')
            report['checks'].append('native_single_header_without_scan_background')
            js('document.querySelector("#composerModel").click()')
            wait_for(lambda: js('document.querySelector("#composerModelChoices button") !== null'))
            assert js('!document.querySelector("#settingsDialog").open')
            js('document.querySelector("#composerModelPopover").dispatchEvent(new KeyboardEvent("keydown", {key:"Escape",bubbles:true}))')
            wait_for(lambda: js('document.querySelector("#composerModelPopover").hidden'))
            assert js('typeof window.pywebview.api.copy_text === "function"')
            report['checks'].append('native_composer_model_picker_and_clipboard_bridge_available')

            def send(text):
                count = js('document.querySelectorAll(".done-event").length')
                js(f'document.querySelector("#prompt").value = {json.dumps(text)}; document.querySelector("#composer").requestSubmit();')
                return count

            def done(count):
                wait_for(lambda: js(f'document.querySelectorAll(".done-event").length > {count} && !document.querySelector("#sendButton").disabled'))
                errors = js('[...document.querySelectorAll(".event.error")].map(e => e.textContent)')
                assert not errors, f'UI reported errors: {errors}'

            stage = 'read'
            done(send('/read sample.txt'))
            assert js('document.querySelector("#prompt").value === ""')
            assert js('document.querySelector(".message.assistant").textContent.includes("original evidence")')
            report['checks'].append('read_file_via_composer')

            stage = 'approve_edit'
            count = send('/edit sample.txt | original | revised')
            wait_for(lambda: js('!!document.querySelector(".approval-actions .primary:not(:disabled)")'))
            assert target.read_text(encoding='utf-8') == 'original evidence\n'
            js('document.querySelector(".approval-actions .primary:not(:disabled)").click()')
            done(count)
            assert target.read_text(encoding='utf-8') == 'revised evidence\n'
            report['checks'].append('approval_gates_write_and_resumes_turn')

            stage = 'reject_edit'
            count = send('/edit sample.txt | revised | rejected')
            wait_for(lambda: js('!!document.querySelector(".approval-actions button:not(:disabled)")'))
            js('document.querySelector(".approval-actions button:not(:disabled)").click()')
            done(count)
            assert target.read_text(encoding='utf-8') == 'revised evidence\n'
            report['checks'].append('rejected_edit_preserves_file')

            stage = 'save_card'
            count = send('/save Desktop routing | Native shell approval evidence | desktop')
            wait_for(lambda: js('!!document.querySelector(".approval-actions .primary:not(:disabled)")'))
            js('document.querySelector(".approval-actions .primary:not(:disabled)").click()')
            done(count)
            assert len(runtime.cards.list_cards()) == 1
            card = runtime.cards.list_cards()[0]
            assert runtime.knowledge.source_summary(card.path)['chunk_count'] == 1
            report['checks'].append('card_approval_and_index')

            stage = 'knowledge_repair'
            source_bytes = Path(card.path).read_bytes()
            with sqlite3.connect(runtime.knowledge.database) as connection:
                connection.execute('DELETE FROM chunks WHERE source_path=?', (card.path,))
            js('document.querySelector("[data-tab=knowledge]").click()')
            wait_for(lambda: js('document.querySelector("#knowledgeStatus").textContent.includes("可修复 1")'))
            js('document.querySelector("#repairKnowledge").click(); document.querySelector("#actionAccept").click();')
            wait_for(lambda: js('document.querySelector("#knowledgeStatus").textContent.includes("一致") && !document.querySelector("#sendButton").disabled'))
            assert Path(card.path).read_bytes() == source_bytes
            assert runtime.knowledge.source_summary(card.path)['chunk_count'] == 1
            report['checks'].append('missing_card_index_repaired_through_native_ui')

            stage = 'provider_diagnostics'
            js('document.querySelector("#settingsButton").click()')
            wait_for(lambda: js('document.querySelector("#settingsDialog").open'))
            js('document.querySelector("#providerDiagnose").click()')
            wait_for(lambda: js('document.querySelectorAll("#providerDiagnostics .diagnostic-row").length === 3 && !document.querySelector("#providerDiagnose").disabled'))
            assert js('document.querySelector("#providerDiagnostics").textContent.includes("离线诊断")')
            assert js('document.querySelectorAll("#settingsForm fieldset legend").length === 3')
            assert js('[...document.querySelectorAll(".icon-button")].every(button => button.querySelector("svg") && button.getAttribute("aria-label"))')
            report['checks'].append('native_grouped_settings_and_accessible_offline_icons')
            assert not runtime.config_repository.path.exists(), 'Diagnostics saved configuration'
            js('document.querySelector("#closeSettings").click()')
            report['checks'].append('mock_capabilities_separated_without_saving_settings')

            stage = 'guidance_settings'
            js('document.querySelector("#settingsButton").click()')
            wait_for(lambda: js('document.querySelector("#settingsDialog").open'))
            assert js('document.querySelector("#guidanceEnabled").checked && document.querySelector("#linkReasoning").disabled')
            js('document.querySelector("#guidanceEditor").open = true; document.querySelector("#guidance-fast-prompt").value = "Native custom guide"; document.querySelector("#guidance-standard-examples").value += "\\nNative custom trigger"; document.querySelector("#settingsForm").requestSubmit()')
            wait_for(lambda: js('!document.querySelector("#settingsDialog").open'))
            assert runtime.config.guidance.profiles.fast.prompt == 'Native custom guide'
            # Additional routing turns can cross the fixture's context threshold.
            runtime.provider = CompactionSmokeProvider()
            done(send('简单回答'))
            js('document.querySelector("#routeDetails").click()')
            assert js('document.querySelector("#routeDetailList").textContent.includes("Native custom guide")')
            js('document.querySelector("#closeRoute").click()')
            report['checks'].append('native_custom_guidance_saved_and_injected')
            done(send('不用深入分析，这次只给结论。'))
            assert js('document.querySelector("#routeDetails").title.includes("explicit_fast")')
            done(send('按正常专业程度回答，解释主要步骤并举例，不超过350字。'))
            assert js('document.querySelector("#routeDetails").title.includes("explicit_standard")')
            report['checks'].append('native_custom_examples_preserve_explicit_intent_and_standard_mode')

            stage = 'nonblocking_card_draft'
            done(send('/draft Reusable finding | Verified native draft | Useful next time'))
            assert len(runtime.cards.list_cards()) == 1
            wait_for(lambda: js('document.querySelectorAll("#draftList button").length === 1'))
            js('document.querySelector("#draftList button").click()')
            wait_for(lambda: js('document.querySelector("#draftDialog").open'))
            js('document.querySelector("#draftContent").value = "Reviewed native finding"; document.querySelector("#draftForm").requestSubmit()')
            wait_for(lambda: len(runtime.cards.list_cards()) == 2)
            wait_for(lambda: js('!document.querySelector("#draftDialog").open'))
            report['checks'].append('native_nonblocking_draft_review_and_confirm')

            stage = 'permission_mode'
            js('document.querySelector("[data-tab=workspace]").click(); document.querySelector("#permissionMode").value = "auto_edit"; document.querySelector("#permissionMode").dispatchEvent(new Event("change")); document.querySelector("#actionAccept").click();')
            wait_for(lambda: runtime.permission_mode == 'auto_edit')
            done(send('/edit sample.txt | revised | automatic'))
            assert target.read_text(encoding='utf-8') == 'automatic evidence\n'
            assert runtime.workspace_audit()
            report['checks'].append('native_opt_in_auto_edit_and_audit')

            stage = 'context_display'
            js('document.querySelector("#contextStatus").click()')
            assert js('document.querySelector("#contextDialog").open && document.querySelector("#contextDetails").textContent.includes("保守")')
            js('document.querySelector("#closeContext").click()')
            assert js('document.querySelector(".brand-mark").naturalWidth > 0')
            report['checks'].append('native_context_details_and_mobius_asset')

            stage = 'native_icon_and_message_attachments'
            from System import Action
            from System.Drawing.Imaging import ImageFormat
            icon_path = root / 'native-window-icon.png'
            window.native.Invoke(Action(lambda: window.native.Icon.ToBitmap().Save(str(icon_path), ImageFormat.Png)))
            assert icon_path.stat().st_size > 100
            from PIL import Image, ImageChops, ImageStat
            from agent_workbench.app import desktop_icon_path
            actual = Image.open(icon_path).convert('RGBA')
            expected = Image.open(desktop_icon_path()).ico.getimage(actual.size).convert('RGBA')
            assert max(ImageStat.Stat(ImageChops.difference(actual, expected)).mean) < 2
            report['checks'].append('native_window_icon_exported')
            from docx import Document
            from openpyxl import Workbook
            from PIL import Image
            word = io.BytesIO(); document = Document(); document.add_paragraph('Native Word evidence'); document.save(word)
            excel = io.BytesIO(); book = Workbook(); book.active.append(['Native Excel evidence', 42]); book.save(excel)
            image = io.BytesIO(); Image.new('RGB', (48, 32), (30, 140, 110)).save(image, format='PNG')
            files = [{'name':name, 'data':base64.b64encode(data).decode('ascii')} for name, data in
                     [('native.docx', word.getvalue()), ('native.xlsx', excel.getvalue()), ('native.png', image.getvalue())]]
            js(f'(() => {{ const transfer=new DataTransfer(); for(const file of {json.dumps(files)}) transfer.items.add(new File([Uint8Array.from(atob(file.data), c=>c.charCodeAt(0))],file.name)); const input=document.querySelector("#attachmentInput"); input.files=transfer.files; input.dispatchEvent(new Event("change",{{bubbles:true}})); }})()')
            wait_for(lambda: len(runtime.attachments.pending()) == 3 and js('!document.querySelector("#sendButton").disabled'))
            done(send('Inspect attached evidence'))
            assert len(runtime.conversation_snapshot()['turns'][-1]['attachments']) == 3
            report['checks'].append('native_word_excel_image_attachments_sent_and_saved')

            stage = 'pdf_import'
            with pymupdf.open() as document:
                page = document.new_page()
                page.insert_text((72, 72), 'Desktop Methods', fontsize=18)
                page.insert_text((72, 110), 'Native desktop routing evidence supports controlled approval workflows.', fontsize=11)
                payload = base64.b64encode(document.tobytes()).decode('ascii')
            js(f'const bytes = Uint8Array.from(atob({json.dumps(payload)}), c => c.charCodeAt(0)); const transfer = new DataTransfer(); transfer.items.add(new File([bytes], "desktop-methods.pdf", {{type:"application/pdf"}})); const input = document.querySelector("#pdfInput"); input.files = transfer.files; input.dispatchEvent(new Event("change", {{bubbles:true}}));')
            wait_for(lambda: runtime.papers.list() and runtime.papers.list()[0].status == 'ready')
            report['checks'].append('pdf_import_through_native_dom_with_real_e5')
            stage = 'paper_tree_and_reader'
            paper = runtime.papers.list()[0]
            runtime.papers.edit_metadata(paper.paper_id, {'title':'Desktop Methods','year':2024,'venue':'Local fixture','topic':'Native verification'}, paper.revision)
            js('document.querySelector("#openPaperLibrary").click()')
            wait_for(lambda: js('document.querySelector("#treeViewport canvas")?.dataset.rendered === "true"'))
            assert js('(() => {const c=document.querySelector("#treeViewport canvas"), gl=c.getContext("webgl2"), bytes=new Uint8Array(gl.drawingBufferWidth*gl.drawingBufferHeight*4);gl.readPixels(0,0,gl.drawingBufferWidth,gl.drawingBufferHeight,gl.RGBA,gl.UNSIGNED_BYTE,bytes);let count=0;for(let i=0;i<bytes.length;i+=4)if(bytes[i]<210||bytes[i+1]<210||bytes[i+2]<210)count++;return count>100;})()')
            js('document.querySelector(".tree-paper:not([hidden])").click()')
            wait_for(lambda: js('!document.querySelector("#paperDetail").hidden'))
            js('document.querySelector("#readPaper").click()')
            wait_for(lambda: js('document.querySelector("#paperPageImage").naturalWidth > 0 && document.querySelector("#paperPageText").textContent.includes("Desktop Methods")'))
            js('document.querySelector("#closePaperReader").click(); document.querySelector("#askPaper").click()')
            assert js('!document.querySelector("#paperScope").hidden && document.querySelector("#paperLibrary").hidden')
            js('document.querySelector("#clearPaperScope").click()')
            report['checks'].append('native_threejs_pixels_paper_reader_and_scoped_question')
            stage = 'knowledge_search'
            # The larger output reserve triggers compaction earlier in this long fixture.
            runtime.provider = CompactionSmokeProvider()
            history_before = {r['id'] for r in runtime.history_store.records()}
            done(send('/kb native desktop routing'))
            assert history_before <= {r['id'] for r in runtime.history_store.records()}
            assert runtime.history_store.checkpoint()
            assert runtime.context_budget(runtime.config).output_tokens == 8192
            report['checks'].append('native_larger_output_reserve_compacts_without_deleting_history')
            assert js('[...document.querySelectorAll(".message.assistant")].at(-1).textContent.includes("Desktop Methods")')
            assert js('[...document.querySelectorAll(".message.assistant")].at(-1).textContent.includes("parser_version")')
            report['checks'].append('knowledge_search_with_section_metadata')

            stage = 'safe_markdown'
            js('window.__awbMarkdownAttack = 0')
            done(send('Render this evidence:\n\n## Native Markdown\n\n**Strong evidence**\n\n| Tool | Result |\n| --- | --- |\n| Read | pass |\n\n```js\nwindow.__awbMarkdownAttack = 1;\n```\n\n<script>window.__awbMarkdownAttack = 2;</script>'))
            assert js('[...document.querySelectorAll(".message.assistant")].at(-1).querySelector("h2").textContent === "Native Markdown"')
            assert js('[...document.querySelectorAll(".message.assistant")].at(-1).querySelectorAll("table tbody td").length === 2')
            assert js('window.__awbMarkdownAttack === 0')
            js('window.__unsafeLinkRejected = false; window.pywebview.api.open_external_url("javascript:alert(1)").catch(() => { window.__unsafeLinkRejected = true; })')
            wait_for(lambda: js('window.__unsafeLinkRejected'))
            report['checks'].append('safe_markdown_and_url_rejection_in_native_shell')
            report['interaction_errors'] = js('window.__awbErrors')
            assert report['interaction_errors'] == []

            stage = 'semantic_compaction'
            runtime.provider = CompactionSmokeProvider()
            runtime.conversation.append('Native compaction fixture', [
                {'role':'user','content':'Native compaction fixture'},
                {'role':'assistant','content':'Verified native desktop evidence. ' * 400}])
            original_count = len(runtime.history_store.records())
            js('document.querySelector("#contextStatus").click(); document.querySelector("#compactNow").click()')
            wait_for(lambda: js('document.querySelector("#compactionStatus").textContent.includes("已完成") && !document.querySelector("#compactNow").disabled'))
            assert len(runtime.history_store.records()) == original_count
            assert runtime.history_store.checkpoint()
            js('document.querySelector("#viewSummary").click()')
            wait_for(lambda: js('document.querySelector("#summaryDialog").open && document.querySelector("#summaryContent").textContent.includes("Keep original history")'))
            js('document.querySelector("#closeSummary").click(); document.querySelector("#closeContext").click()')
            report['checks'].append('native_semantic_compaction_with_fixture_provider')

            stage = 'history_reload'
            js('location.reload()')
            expected_turns = len(runtime.conversation.turns)
            wait_for(lambda: js(f'document.querySelectorAll(".message.user").length === {expected_turns}'))
            report['checks'].append('history_restores_after_reload')
            stage = 'resize'
            report['viewports'] = []
            for width, height in [(1366, 768), (1920, 1080), (960, 600)]:
                window.resize(width, height)
                time.sleep(0.5)
                metrics = js('({width: innerWidth, height: innerHeight, scroll: document.documentElement.scrollWidth, client: document.documentElement.clientWidth})')
                assert metrics['scroll'] <= metrics['client'], metrics
                assert js('(() => { const a = document.querySelector(".message.assistant").getBoundingClientRect(); const b = document.querySelector(".composer-inner").getBoundingClientRect(); return Math.abs(a.left-b.left) <= 1 && Math.abs(a.right-b.right) <= 1; })()')
                report['viewports'].append(metrics)
            report['checks'].append('native_resize_without_document_horizontal_overflow')
            stage = 'history_clear'
            js('document.querySelector("#clearConversation").click(); document.querySelector("#actionAccept").click();')
            wait_for(lambda: not runtime.conversation_snapshot()['turns'])
            wait_for(lambda: js('document.querySelectorAll(".message.user").length === 0'))
            report['checks'].append('clear_history_synchronizes')
            report['status'] = 'passed'
        except Exception:
            report['status'] = 'failed'
            report['stage'] = stage
            report['error'] = traceback.format_exc()
        finally:
            if window is not None:
                if report.get('status') == 'passed' and hasattr(window, '_desktop_tray'):
                    native_call(lambda: tray.exit_item.PerformClick())
                    report['checks'].append('native_tray_exit_menu_closes_application')
                else:
                    if hasattr(window, '_desktop_tray'):
                        window._desktop_tray.controller.exiting = True
                    window.destroy()

    worker = threading.Thread(target=exercise, daemon=True)
    worker.start()
    try:
        run_desktop(runtime, port)
        worker.join(timeout=5)
        with socket.socket() as connection:
            assert connection.connect_ex(('127.0.0.1', port)) != 0, 'backend remains listening after window close'
        assert not any(thread.name == 'awb-api' for thread in threading.enumerate())
        report['checks'].append('window_close_stops_backend')
    except Exception:
        report['status'] = 'failed'
        report['error'] = traceback.format_exc()
    report_path = root / 'report.json'
    report_path.write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(json.dumps(report, indent=2), flush=True)
    print(f'Report: {report_path}', flush=True)
    return 0 if report['status'] == 'passed' else 1


if __name__ == '__main__':
    raise SystemExit(main())
