"""Exercise the packaged EXE with isolated data and a loopback-only model fixture."""
import json
import io
import os
from pathlib import Path
import subprocess
import tempfile
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import httpx

from agent_workbench.app import find_available_port
from agent_workbench import __version__
from agent_workbench.config import AppConfig, ConfigRepository
from agent_workbench.core.history import HistoryStore


def main():
    project = Path(__file__).resolve().parents[1]
    root = Path(tempfile.mkdtemp(prefix='packaged-v15-', dir=project / '.tmp'))
    requests = []

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def do_POST(self):
            # Never log credentials, and reject accidental use of anything except the fixture key.
            if self.headers.get('Authorization') != 'Bearer awb-local-fixture':
                self.send_error(401)
                return
            payload = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
            requests.append(payload)
            compact = 'conversation compaction' in payload['messages'][0]['content']
            content = '## Constraints\nKeep original history.\n## Progress\nVerified SQLite.\n## Pending\nCheck upgrade.' if compact else 'Packaged continuation verified'
            if 'Classify one research paper.' in payload['messages'][0]['content']:
                content = json.dumps({'title':'Packaged Paper','year':2024,'year_type':'published','venue':'Fixture Journal','topic':'Verification',
                                      'evidence':{'title':{'page':1,'quote':'Packaged Paper'},'year':{'page':1,'quote':'2024'},'venue':{'page':1,'quote':'Fixture Journal'}}})
            body = ('data: ' + json.dumps({'choices':[{'delta':{'content':content},'finish_reason':'stop'}]})
                    + '\n\ndata: ' + json.dumps({'choices':[],'usage':{'prompt_tokens':100,'completion_tokens':20,'total_tokens':120}})
                    + '\n\ndata: [DONE]\n\n').encode()
            self.send_response(200)
            self.send_header('Content-Type', 'text/event-stream')
            self.send_header('Content-Length', str(len(body)))
            self.end_headers()
            self.wfile.write(body)

    upstream = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
    thread = threading.Thread(target=upstream.serve_forever, daemon=True)
    thread.start()
    config = AppConfig(provider_mode='openai-compatible', base_url=f'http://127.0.0.1:{upstream.server_port}/v1',
                       model='deepseek-v4-flash', reasoning_effort='high', context_window=32768, max_output_tokens=1024)
    config.guidance.enabled = True
    config.guidance.profiles.standard.examples.append('Packaged custom trigger')
    ConfigRepository(root / 'config.json').save(config)
    store = HistoryStore(root / 'history.sqlite3')
    key = json.dumps([config.provider_mode,config.base_url,config.model])
    for i in range(8):
        store.append(str(i), [{'role':'user','content':f'Fixture request {i}'},
                             {'role':'assistant','content':'Verified SQLite evidence. ' * 150}], 'stop', provider_key=key)
    # Reproduce the V11 schema so the EXE must add its new tables without losing records.
    with store.connect() as db:
        db.execute('DROP TABLE checkpoints')
        db.execute('DROP TABLE compaction_calls')
    environment = {**os.environ, 'OPENAI_API_KEY':'awb-local-fixture',
                   'PYTHON_KEYRING_BACKEND':'keyring.backends.null.Keyring',
                   'TEMP':str(root), 'TMP':str(root), 'NO_PROXY':'127.0.0.1,localhost'}
    executable = project / 'dist' / 'AgentWorkbench' / 'AgentWorkbench.exe'
    child = None
    report = {'status':'running', 'data_dir':str(root), 'checks':[]}
    log = (root / 'process.log').open('ab')

    def start():
        nonlocal child
        port = find_available_port()
        child = subprocess.Popen([str(executable), '--server-only', '--port',str(port),'--data-dir',str(root)],
                                 env=environment, stdout=log, stderr=log, creationflags=subprocess.CREATE_NO_WINDOW)
        url = f'http://127.0.0.1:{port}'
        deadline = time.monotonic() + 90
        while time.monotonic() < deadline:
            if child.poll() is not None:
                raise RuntimeError('Packaged process exited during startup; inspect process.log')
            try:
                if httpx.get(url + '/health/live', timeout=1, trust_env=False).status_code == 200:
                    return url
            except httpx.HTTPError:
                pass
            time.sleep(.2)
        raise TimeoutError('Packaged startup timed out')

    def stop():
        nonlocal child
        if child is not None:
            child.terminate()
            try:
                child.wait(timeout=10)
            except subprocess.TimeoutExpired:
                child.kill()
                child.wait(timeout=10)
            child = None

    try:
        url = start()
        with httpx.Client(base_url=url, timeout=90, trust_env=False) as client:
            assert client.get('/openapi.json').json()['info']['version'] == __version__
            limits = client.get('/api/settings').json()
            assert limits['max_output_tokens'] == 1024
            assert limits['effective_limits'] == {'context_window':32768,'max_output_tokens':1024}
            assert 'ambient-plane' not in client.get('/').text
            assert client.get('/assets/app.css').content == (project / 'assets' / 'web' / 'app.css').read_bytes()
            for name in ('paper-tree.js','paper-library.js','paper-library.css','dialogs.js','model-picker.js','composer-model.js','clipboard.js','markdown.js','app.js','desktop-features.js','ambient.js','mobius-scene.js'):
                assert client.get('/assets/' + name).content == (project / 'assets' / 'web' / name).read_bytes()
            report['checks'].append('packaged_frontend_matches_single_header_ambient_release')
            before = client.get('/api/conversation').json()
            assert len(before['turns']) == 8
            result = client.post('/api/context/compact')
            result.raise_for_status()
            assert result.json()['status'] == 'completed'
            checkpoint = client.get('/api/context/summary').json()['checkpoint']
            assert checkpoint and len(client.get('/api/conversation').json()['turns']) == 8
            report['checks'].append('v11_schema_upgrade_and_manual_semantic_checkpoint')
            message = '按正常专业程度回答，不超过350字。Continue the task'
            response = client.post('/api/chat/stream',json={'message':message})
            assert 'Packaged continuation verified' in response.text
            assert 'explicit_standard' in response.text
            assert requests[-1]['messages'][-1]['content'] == message
            assert any('<compacted-summary>' in (m.get('content') or '') for m in requests[-1]['messages'])
            assert all(p['reasoning_effort'] == 'high' for p in requests)
            # Compaction has its own bounded output budget; the chat keeps the user's value.
            assert requests[-1]['max_tokens'] == 1024
            report['checks'].append('packaged_provider_continues_with_summary_and_manual_reasoning')
        stop()
        url = start()
        with httpx.Client(base_url=url, timeout=90, trust_env=False) as client:
            restored = client.get('/api/conversation').json()
            assert restored['session_id'] == before['session_id'] and len(restored['turns']) == 9
            assert client.get('/api/context/summary').json()['checkpoint'] == checkpoint
            assert client.post('/api/context/estimate',json={'message':''}).json()['summary_kind'] == 'semantic'
            report['checks'].append('checkpoint_and_complete_history_survive_packaged_restart')
        stop()
        url = start()
        with httpx.Client(base_url=url, timeout=90, trust_env=False) as client:
            assert client.get('/api/context/summary').json()['checkpoint'] == checkpoint
            report['checks'].append('second_packaged_restart_preserves_checkpoint')
            from docx import Document
            from openpyxl import Workbook
            from PIL import Image
            word = io.BytesIO()
            document = Document(); document.add_paragraph('Packaged Word evidence'); document.save(word)
            excel = io.BytesIO()
            book = Workbook(); book.active.append(['Packaged Excel evidence', 42]); book.save(excel)
            image = io.BytesIO()
            Image.new('RGB', (64, 48), (30, 140, 110)).save(image, format='PNG')
            attachments = []
            legacy = (project / 'tests' / 'fixtures' / 'attachment.xls').read_bytes()
            for name, data in [('report.docx', word.getvalue()), ('data.xlsx', excel.getvalue()), ('legacy.xls', legacy), ('image.png', image.getvalue())]:
                result = client.post('/api/attachments', params={'filename':name}, content=data,
                                     headers={'content-type':'application/octet-stream'})
                result.raise_for_status()
                attachments.append((result.json()['id'], data))
            result = client.post('/api/chat/stream', json={'message':'Inspect the files', 'attachment_ids':[key for key, _ in attachments]})
            assert 'Packaged continuation verified' in result.text
            blocks = requests[-1]['messages'][-1]['content']
            assert isinstance(blocks, list) and any(p['type'] == 'image_url' for p in blocks)
            assert 'Packaged Word evidence' in json.dumps(blocks) and 'Packaged Excel evidence' in json.dumps(blocks)
            assert 'Legacy Excel evidence' in json.dumps(blocks)
            assert 'base64' not in client.get('/api/conversation').text
            report['checks'].append('packaged_word_excel_image_extraction_and_multimodal_provider_payload')
        stop()
        url = start()
        with httpx.Client(base_url=url, timeout=90, trust_env=False) as client:
            for key, data in attachments:
                assert client.get(f'/api/attachments/{key}/content').content == data
            result = client.post('/api/chat/stream', json={'message':'Continue reviewing'})
            assert 'Packaged continuation verified' in result.text
            assert any(isinstance(m.get('content'), list) and any(p.get('type') == 'image_url' for p in m['content']) for m in requests[-1]['messages'])
            report['checks'].append('packaged_restart_retains_original_attachments_and_replays_image')
            settings = client.get('/api/settings').json()
            assert settings['dynamic_background'] is False
            settings['dynamic_background'] = True
            client.put('/api/settings', json=settings).raise_for_status()
        stop()
        url = start()
        with httpx.Client(base_url=url, timeout=90, trust_env=False) as client:
            assert client.get('/api/settings').json()['dynamic_background'] is True
            assert 'dynamicBackground' not in client.get('/').text
            report['checks'].append('legacy_appearance_field_survives_restart_without_restoring_background_ui')
            import pymupdf
            with pymupdf.open() as pdf:
                pdf.new_page().insert_text((72,72),'Packaged Paper\n2024\nFixture Journal')
                paper_data=pdf.tobytes()
            response=client.post('/api/kb/papers',content=paper_data,headers={'content-type':'application/pdf','x-filename':'fixture.pdf'})
            response.raise_for_status();paper_id=response.json()['paper_id']
            client.post(f'/api/kb/papers/{paper_id}/classify?confirmed=true').raise_for_status()
            deadline=time.monotonic()+60
            while time.monotonic()<deadline:
                paper=client.get(f'/api/kb/papers/{paper_id}').json()
                if paper['classification_status'] not in {'queued','classifying'}:break
                time.sleep(.15)
            assert paper['classification_status']=='ready',paper.get('classification_error')
            assert paper['metadata']['year']==2024 and paper['topic_slot']==0
            assert client.get(f'/api/kb/papers/{paper_id}/pages/1?image=true').content.startswith(b'\x89PNG')
            assert client.get('/assets/vendor/three/three.module.js').status_code==200
            report['checks'].append('packaged_pdf_classification_grounding_reader_and_offline_threejs')
        stop()
        url=start()
        with httpx.Client(base_url=url,timeout=90,trust_env=False) as client:
            assert client.get(f'/api/kb/papers/{paper_id}').json()['metadata']==paper['metadata']
            assert client.get(f'/api/kb/papers/{paper_id}/content').content==paper_data
            report['checks'].append('packaged_paper_metadata_topic_layout_and_original_survive_restart')
        report['status'] = 'passed'
        report['model_fixture_requests'] = len(requests)
    except Exception as exc:
        report['status'] = 'failed'
        report['error'] = type(exc).__name__
        raise
    finally:
        stop()
        upstream.shutdown()
        upstream.server_close()
        thread.join(timeout=5)
        log.close()
        (root / 'report.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps(report,indent=2))


if __name__ == '__main__':
    main()
