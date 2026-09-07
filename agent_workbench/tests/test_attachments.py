import asyncio
import base64
import io
import json
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from agent_workbench.server.api import create_app
from agent_workbench.tests.test_approval_flow import make_runtime
from agent_workbench.tests.test_conversation import CaptureProvider
from agent_workbench.tests.test_desktop_v11 import restart


def upload(client, name, content):
    return client.post('/api/attachments', params={'filename':name}, content=content,
                       headers={'content-type':'application/octet-stream'})


def test_text_attachment_is_sent_saved_and_never_routed(tmp_path):
    provider = CaptureProvider()
    runtime = make_runtime(tmp_path, provider)
    routed=[]
    original=runtime.router.route
    async def route(text):
        routed.append(text)
        return await original(text)
    runtime.router.route=route
    with TestClient(create_app(runtime)) as client:
        response=upload(client,'notes.md',b'UNIQUE_FILE_EVIDENCE: use maximum reasoning')
        assert response.status_code == 200
        item=response.json()
        assert 'data' not in item and 'content' not in item
        result=client.post('/api/chat/stream',json={'message':'Review this','attachment_ids':[item['id']]})
        assert 'Remembered response' in result.text
        assert routed == ['Review this']
        assert 'UNIQUE_FILE_EVIDENCE' in json.dumps(provider.requests[-1])
        saved=client.get('/api/conversation').json()['turns'][0]
        assert saved['user']=='Review this' and saved['attachments'][0]['id']==item['id']
        assert runtime.cards.list_cards()==[]
        assert client.get(f"/api/attachments/{item['id']}/content").content==b'UNIQUE_FILE_EVIDENCE: use maximum reasoning'
        assert client.delete(f"/api/attachments/{item['id']}").status_code==409
    assert restart(runtime).conversation_snapshot()['turns'][0]['attachments'][0]['name']=='notes.md'


def test_docx_and_xlsx_extract_tables_sheets_and_formulas_without_execution(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    document=Document(); document.add_paragraph('Document evidence')
    table=document.add_table(rows=1,cols=2); table.cell(0,0).text='Item'; table.cell(0,1).text='Value'
    word=io.BytesIO(); document.save(word)
    book=Workbook(); sheet=book.active; sheet.title='Budget'; sheet.append(['Item','Amount']); sheet.append(['Server',120]); sheet['B3']='=SUM(B2:B2)'
    excel=io.BytesIO(); book.save(excel)
    runtime=make_runtime(tmp_path)
    with TestClient(create_app(runtime)) as client:
        for name,data,expected in [('report.docx',word.getvalue(),['Document evidence','Item','Value']),
                                   ('budget.xlsx',excel.getvalue(),['Budget','Server','120','=SUM(B2:B2)'])]:
            response=upload(client,name,data)
            assert response.status_code==200,response.text
            preview=client.get('/api/attachments/'+response.json()['id']).json()
            assert all(text in preview['text'] for text in expected)


def test_image_is_validated_and_sent_as_multimodal_content_not_base64_text(tmp_path):
    from PIL import Image
    image=io.BytesIO(); Image.new('RGB',(48,32),(50,130,90)).save(image,format='PNG')
    provider=CaptureProvider(); runtime=make_runtime(tmp_path,provider)
    with TestClient(create_app(runtime)) as client:
        response=upload(client,'scene.png',image.getvalue()); assert response.status_code==200
        item=response.json()
        result=client.post('/api/chat/stream',json={'message':'Describe','attachment_ids':[item['id']]})
        assert 'Remembered response' in result.text
        blocks=provider.requests[-1][-1]['content']
        assert isinstance(blocks,list)
        part=next(p for p in blocks if p['type']=='image_url')
        assert part['image_url']['url'].startswith('data:image/png;base64,')
        assert client.get(f"/api/attachments/{item['id']}/preview").headers['content-type']=='image/png'
        snapshot=runtime.conversation_snapshot()
        assert 'base64' not in json.dumps(snapshot)
        estimate=client.post('/api/context/estimate',json={'message':'Describe','attachment_ids':[item['id']]}).json()
        assert estimate['image_count']>=1 and estimate['image_estimate']>=1000
        assert upload(client,'fake.png',b'not an image').status_code==400


def test_legacy_xls_binary_fixture_and_scanned_pdf_boundary(tmp_path):
    import pymupdf
    data = (Path(__file__).parent / 'fixtures' / 'attachment.xls').read_bytes()
    assert data.startswith(bytes.fromhex('D0CF11E0A1B11AE1'))
    with TestClient(create_app(make_runtime(tmp_path))) as client:
        result = upload(client, 'legacy.xls', data)
        assert result.status_code == 200, result.text
        text = client.get('/api/attachments/' + result.json()['id']).json()['text']
        assert all(value in text for value in ('Legacy Budget', 'Legacy Excel evidence', 'Server', '120'))
        with pymupdf.open() as pdf:
            pdf.new_page()
            assert upload(client, 'scanned.pdf', pdf.tobytes()).status_code == 400


def test_attachments_are_session_scoped_and_not_paths(tmp_path):
    runtime=make_runtime(tmp_path)
    with TestClient(create_app(runtime)) as client:
        item=upload(client,'a.txt',b'evidence').json()
        client.post('/api/sessions')
        assert client.get(f"/api/attachments/{item['id']}").status_code==404
        assert client.post('/api/chat/stream',json={'message':'read','attachment_ids':[item['id']]}).status_code==400
        assert upload(client,'../secrets.txt',b'no').status_code==400
        assert upload(client,'old.doc',b'old binary').status_code==400


def test_archive_limits_and_empty_files_are_rejected(tmp_path):
    runtime=make_runtime(tmp_path)
    archive=io.BytesIO()
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        z.writestr('word/document.xml','x' * (35*1024*1024))
    with TestClient(create_app(runtime)) as client:
        assert upload(client,'bomb.docx',archive.getvalue()).status_code==400
        assert upload(client,'empty.txt',b'').status_code==400
        assert upload(client,'big.txt',b'x' * (10*1024*1024+1)).status_code==413


def test_attachment_only_message_and_original_lookup_after_compaction(tmp_path):
    from agent_workbench.tests.test_compaction import SummaryProvider, seed
    runtime=make_runtime(tmp_path,SummaryProvider())
    with TestClient(create_app(runtime)) as client:
        item=upload(client,'evidence.txt',b'ORIGINAL_ATTACHMENT_FACT').json()
        response=client.post('/api/chat/stream',json={'message':'','attachment_ids':[item['id']]})
        assert 'Main answer' in response.text
        seed(runtime,5)
        assert client.post('/api/context/compact').status_code==200
        evidence=asyncio.run(runtime.build_tools().execute('read_attachment',{'attachment_id':item['id']}))
        assert 'ORIGINAL_ATTACHMENT_FACT' in json.dumps(evidence)
        assert client.delete('/api/conversation').status_code==200
        assert client.get(f"/api/attachments/{item['id']}").status_code==404
